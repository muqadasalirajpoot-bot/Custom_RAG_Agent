from langchain_openai import ChatOpenAI
from langchain_core.messages import SystemMessage, HumanMessage, AIMessage
import streamlit as st
import base64
import io

st.set_page_config(
    page_title="Your AI Assistant!",
    page_icon="🤖",
    layout="centered"
)

# Theme state
if "theme" not in st.session_state:
    st.session_state.theme = "light"

_THEMES = {
    "light": {
        "app_bg": "linear-gradient(180deg, #f7f8fc 0%, #ffffff 35%)",
        "text": "#20202b",
        "sidebar_bg": "#f4f5fa",
        "sidebar_border": "#e6e6f0",
        "bubble_bg": "#ffffff",
        "bubble_shadow": "rgba(0,0,0,0.06)",
        "chip_bg": "#eef0fb",
        "chip_text": "#4b4b6b",
        "subtitle": "#8a8aa3",
        "button_bg": "#ffffff",
        "button_border": "#d8d8e6",
    },
    "dark": {
        "app_bg": "linear-gradient(180deg, #14141c 0%, #1c1c26 35%)",
        "text": "#eaeaf2",
        "sidebar_bg": "#1a1a24",
        "sidebar_border": "#2c2c3a",
        "bubble_bg": "#23232f",
        "bubble_shadow": "rgba(0,0,0,0.35)",
        "chip_bg": "#2a2a3a",
        "chip_text": "#c7c7e0",
        "subtitle": "#9a9ab0",
        "button_bg": "#23232f",
        "button_border": "#3a3a4a",
    },
}
_T = _THEMES[st.session_state.theme]

# --- CUSTOM STYLING ---
st.markdown(f"""
<style>
    /* Page background */
    .stApp {{
        background: {_T['app_bg']};
    }}
    /* Text color: headers, labels, captions, markdown, chat bubbles —
       NOT the actual input/textarea/select boxes, which stay white-bg
       with their own default (readable) text color regardless of theme */
    .stApp h1, .stApp h2, .stApp h3, .stApp h4, .stApp h5, .stApp h6,
    .stMarkdown p, .stMarkdown li, .stMarkdown span,
    [data-testid="stCaptionContainer"], [data-testid="stCaptionContainer"] p,
    [data-testid="stWidgetLabel"] p,
    [data-testid="stChatMessage"] p, [data-testid="stChatMessage"] span,
    section[data-testid="stSidebar"] h1, section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3, section[data-testid="stSidebar"] label,
    section[data-testid="stSidebar"] p, section[data-testid="stSidebar"] span {{
        color: {_T['text']} !important;
    }}
    /* Chat bubbles */
    [data-testid="stChatMessage"] {{
        background: {_T['bubble_bg']} !important;
        border-radius: 16px;
        padding: 4px 6px;
        margin-bottom: 10px;
        box-shadow: 0 1px 3px {_T['bubble_shadow']};
    }}
    /* Buttons (suggestion chips, clear/download buttons) */
    .stApp button[kind="secondary"], .stApp .stButton button {{
        background: {_T['button_bg']} !important;
        color: {_T['text']} !important;
        border: 1px solid {_T['button_border']} !important;
    }}
    .stApp .stDownloadButton button {{
        background: {_T['button_bg']} !important;
        color: {_T['text']} !important;
        border: 1px solid {_T['button_border']} !important;
    }}
    /* Sidebar */
    section[data-testid="stSidebar"] {{
        background: {_T['sidebar_bg']};
        border-right: 1px solid {_T['sidebar_border']};
    }}
    /* Chat input bar */
    [data-testid="stChatInput"] {{
        border-radius: 14px;
    }}
    /* Header gradient text */
    .hero-title {{
        background: linear-gradient(90deg, #6a5cff, #ff5ca7);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        font-weight: 800;
        font-size: 2rem;
        margin-bottom: 0.1rem;
    }}
    .hero-subtitle {{
        color: {_T['subtitle']};
        font-size: 0.95rem;
        margin-bottom: 1.4rem;
    }}
    /* Alerts (st.info) and captions */
    .stApp [data-testid="stAlertContainer"] {{
        background: {_T['bubble_bg']} !important;
    }}
    .stApp [data-testid="stAlertContainer"] p {{
        color: {_T['text']} !important;
    }}
    .attach-chip {{
        display: inline-block;
        background: {_T['chip_bg']};
        color: {_T['chip_text']};
        border-radius: 999px;
        padding: 3px 12px;
        font-size: 0.82rem;
        margin: 2px 4px 2px 0;
    }}
</style>
""", unsafe_allow_html=True)

USER_AVATAR = "🧑"
ASSISTANT_AVATAR = "🤖"
FILE_ICONS = {
    "pdf": "📕", "docx": "📄", "xlsx": "📊", "xls": "📊",
    "csv": "📈", "txt": "📝", "md": "📝", "json": "🗂️", "py": "🐍", "log": "🧾"
}

# Initialize message history in session state
if 'messages' not in st.session_state:
    st.session_state.messages = []

# Sidebar for API Key and Configuration
with st.sidebar:
    theme_dark = st.toggle("🌙 Dark theme", value=(st.session_state.theme == "dark"))
    new_theme = "dark" if theme_dark else "light"
    if new_theme != st.session_state.theme:
        st.session_state.theme = new_theme
        st.rerun()

    st.markdown("## 🔐 Authentication")
    api_key = st.text_input(label="OpenAI API Key", type="password", help="Your key stays local to this session.")

    st.divider()
    st.markdown("## ⚙️ Model")
    model_choice = st.selectbox(
        "Model",
        ["gpt-4o-mini", "gpt-4o", "gpt-3.5-turbo"],
        index=0,
        help="gpt-3.5-turbo cannot see attached images."
    )
    temperature = st.slider("Creativity (temperature)", 0.0, 1.0, 0.5, 0.1)

    st.divider()
    if st.button("🗑️ Clear conversation", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

    # Stats + export
    _human_count = sum(1 for m in st.session_state.messages if isinstance(m, HumanMessage))
    _ai_count = sum(1 for m in st.session_state.messages if isinstance(m, AIMessage))
    if _human_count or _ai_count:
        st.caption(f"💬 {_human_count} sent · {_ai_count} received")

        def _msg_text(m):
            if isinstance(m.content, list):
                return next((p["text"] for p in m.content if p.get("type") == "text"), "")
            return m.content

        _transcript = "\n\n".join(
            f"{'You' if isinstance(m, HumanMessage) else 'Assistant'}: {_msg_text(m)}"
            for m in st.session_state.messages if isinstance(m, (HumanMessage, AIMessage))
        )
        st.download_button(
            "⬇️ Download chat",
            data=_transcript,
            file_name="conversation.txt",
            use_container_width=True
        )

# 1. HIDE ENTIRE UI UNTIL API KEY IS ENTERED
if not api_key:
    st.markdown('<div class="hero-title">Your Custom ChatGPT</div>', unsafe_allow_html=True)
    st.info("👈 Please enter your OpenAI API key in the sidebar to unlock the chatbot.")
    st.stop()

# --- Everything below this line ONLY shows after the user enters their API Key ---

st.markdown('<div class="hero-title">Your Custom ChatGPT</div>', unsafe_allow_html=True)

_has_chatted = any(isinstance(m, HumanMessage) for m in st.session_state.messages)
if not _has_chatted:
    st.markdown('<div class="hero-subtitle">Try one of these, or ask anything below.</div>', unsafe_allow_html=True)
    _suggestions = [
        "💡 Explain this simply",
        "🐍 Write a Python script",
        "📄 Summarize an attached file",
    ]
    _cols = st.columns(len(_suggestions))
    for _col, _s in zip(_cols, _suggestions):
        if _col.button(_s, use_container_width=True):
            st.session_state.pending_prompt = _s.split(" ", 1)[1]
            st.rerun()

# Initialize ChatOpenAI instance with the provided API key
# NOTE: gpt-3.5-turbo has no vision support — pick gpt-4o / gpt-4o-mini in the
# sidebar if you want the model to actually understand attached images.
chat = ChatOpenAI(model_name=model_choice, temperature=temperature, api_key=api_key)

# Display conversation history on the main page
for msg in st.session_state.messages:
    if isinstance(msg, HumanMessage):
        with st.chat_message("user", avatar=USER_AVATAR):
            if isinstance(msg.content, list):
                for part in msg.content:
                    if part.get("type") == "text":
                        st.write(part["text"])
                    elif part.get("type") == "image_url":
                        st.image(part["image_url"]["url"], width=220)
            else:
                st.write(msg.content)
    elif isinstance(msg, AIMessage):
        with st.chat_message("assistant", avatar=ASSISTANT_AVATAR):
            st.write(msg.content)

# --- FILE / IMAGE ATTACHMENT SUPPORT ---
IMAGE_TYPES = {"png", "jpg", "jpeg", "gif", "webp"}


def extract_text_from_file(uploaded_file):
    """Best-effort text extraction for non-image files, so their content
    can be folded into the prompt as extra context."""
    name = uploaded_file.name
    ext = name.split(".")[-1].lower()
    raw = uploaded_file.read()
    uploaded_file.seek(0)

    try:
        if ext in ("txt", "md", "csv", "json", "py", "log"):
            return raw.decode("utf-8", errors="ignore")
        elif ext == "pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return "[Could not extract PDF text: pypdf not installed]"
            reader = PdfReader(io.BytesIO(raw))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == "docx":
            try:
                import docx
            except ImportError:
                return "[Could not extract DOCX text: python-docx not installed]"
            document = docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in document.paragraphs)
        elif ext in ("xlsx", "xls"):
            try:
                import openpyxl
            except ImportError:
                return "[Could not extract Excel text: openpyxl not installed]"
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    lines.append(", ".join(str(c) for c in row if c is not None))
            return "\n".join(lines)
        else:
            return f"[Unsupported file type '{ext}': could not extract text, filename attached only]"
    except Exception as e:
        return f"[Error reading {name}: {e}]"


# 2. CENTER PROMPT BAR AT THE BOTTOM OF MAIN INTERFACE
# accept_file="multiple" adds a paperclip/attach button directly inside the
# chat input bar, so users can attach files or images right where they type.
chat_value = st.chat_input(
    "Ask something...",
    accept_file="multiple",
    file_type=None
)

if chat_value or st.session_state.get("pending_prompt"):
    if chat_value:
        user_prompt = chat_value.text
        uploaded_files = chat_value.files
    else:
        user_prompt = st.session_state.pop("pending_prompt")
        uploaded_files = None

    # Build the message content: plain text if no attachments, otherwise a
    # multimodal list (text + inline images + extracted text from other files)
    final_content = user_prompt
    if uploaded_files:
        content_parts = [{"type": "text", "text": user_prompt}]
        file_context_snippets = []

        for f in uploaded_files:
            ext = f.name.split(".")[-1].lower()
            if ext in IMAGE_TYPES:
                img_bytes = f.read()
                f.seek(0)
                b64 = base64.b64encode(img_bytes).decode("utf-8")
                mime = f"image/{'jpeg' if ext == 'jpg' else ext}"
                content_parts.append({
                    "type": "image_url",
                    "image_url": {"url": f"data:{mime};base64,{b64}"}
                })
            else:
                extracted = extract_text_from_file(f)
                file_context_snippets.append(f"--- Content of {f.name} ---\n{extracted}")

        if file_context_snippets:
            content_parts.append({
                "type": "text",
                "text": "\n\n".join(file_context_snippets)
            })

        final_content = content_parts

    # Append user prompt and render immediately
    st.session_state.messages.append(HumanMessage(content=final_content))
    with st.chat_message("user", avatar=USER_AVATAR):
        if user_prompt:
            st.write(user_prompt)
        if uploaded_files:
            chips = ""
            for f in uploaded_files:
                ext = f.name.split(".")[-1].lower()
                if ext not in IMAGE_TYPES:
                    icon = FILE_ICONS.get(ext, "📎")
                    chips += f'<span class="attach-chip">{icon} {f.name}</span>'
            if chips:
                st.markdown(chips, unsafe_allow_html=True)
            for f in uploaded_files:
                ext = f.name.split(".")[-1].lower()
                if ext in IMAGE_TYPES:
                    st.image(f, width=220)

    # Generate model response, streamed token-by-token for a livelier feel
    def _stream_reply():
        for chunk in chat.stream(st.session_state.messages):
            if chunk.content:
                yield chunk.content

    with st.chat_message("assistant", avatar=ASSISTANT_AVATAR):
        full_response = st.write_stream(_stream_reply())
    st.session_state.messages.append(AIMessage(content=full_response))